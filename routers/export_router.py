"""Xuất báo cáo lỗi DQ ra Excel (XLSX) / CSV theo đúng bộ lọc hiện tại."""
import io
import logging
from datetime import datetime

import pandas as pd
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from auth import get_current_user
from config import settings
from database import get_db
from models import EXPORT_COLUMNS, CurrentUser, DQErrorOut, DQFilter, Severity
from repository import fetch_errors
from routers.dq_router import get_filters
from workflow import audit

logger = logging.getLogger("dq.export")
router = APIRouter(prefix="/export", tags=["export"])

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SEVERITY_FILL = {
    Severity.HIGH.value: PatternFill("solid", fgColor="F8CBAD"),
    Severity.MEDIUM.value: PatternFill("solid", fgColor="FFE699"),
    Severity.LOW.value: PatternFill("solid", fgColor="DDEBF7"),
}
COLUMN_WIDTHS = {"STT": 7, "branch_code": 13, "branch_name": 28, "rule_code": 15, "description": 60,
                 "severity": 10, "detected_date": 15, "status": 20, "owner": 22, "resolved_date": 17,
                 "resolved_by": 18}
DATE_COLUMNS = ("detected_date", "resolved_date")


def _scope_label(f: DQFilter, user: CurrentUser) -> str:
    if user.all_branches and not f.branch_code:
        return "Toàn hệ thống"
    return f"Chi nhánh {f.branch_code}"


def _describe_filters(f: DQFilter) -> str:
    parts = []
    if f.severity:
        parts.append(f"Mức độ: {f.severity.value}")
    if f.status:
        parts.append(f"Trạng thái: {f.status.value}")
    if f.keyword:
        parts.append(f"Từ khóa: \"{f.keyword}\"")
    if f.date_from or f.date_to:
        start = f.date_from.strftime("%d/%m/%Y") if f.date_from else "..."
        end = f.date_to.strftime("%d/%m/%Y") if f.date_to else "..."
        parts.append(f"Ngày phát hiện: {start} - {end}")
    return "; ".join(parts) or "Không có"


def _filename(ext: str, f: DQFilter, user: CurrentUser) -> str:
    scope = "ToanHeThong" if user.all_branches and not f.branch_code else f.branch_code
    return f"DQ_Report_{scope}_{datetime.now():%Y%m%d_%H%M%S}.{ext}"


def _load_rows(db: Session, f: DQFilter, user: CurrentUser, fmt: str) -> tuple[list[DQErrorOut], bool]:
    """Lấy limit+1 bản ghi để biết dữ liệu có bị cắt bớt hay không; ghi nhật ký xuất file."""
    rows = fetch_errors(db, f, settings.EXPORT_MAX_ROWS + 1)
    truncated = len(rows) > settings.EXPORT_MAX_ROWS
    count = min(len(rows), settings.EXPORT_MAX_ROWS)
    audit(db, user.username, "EXPORT", fmt.upper(),
          f"{_scope_label(f, user)}; bộ lọc: {_describe_filters(f)}; {count} dòng")
    db.commit()
    logger.info("Export %s bởi %s: phạm vi=%s, %d dòng", fmt, user.username, _scope_label(f, user), count)
    return rows[:settings.EXPORT_MAX_ROWS], truncated


def build_excel(rows: list[DQErrorOut], user: CurrentUser, f: DQFilter, truncated: bool) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Chi tiết lỗi DQ"
    _write_detail_sheet(ws, rows, user, f, truncated)
    _write_summary_sheet(wb.create_sheet("Tổng hợp"), rows)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _write_detail_sheet(ws, rows: list[DQErrorOut], user: CurrentUser, f: DQFilter, truncated: bool) -> None:
    """Sheet chi tiết: tiêu đề, thông tin bộ lọc, header và dữ liệu."""
    keys = ["STT", *EXPORT_COLUMNS.keys()]
    headers = ["STT", *EXPORT_COLUMNS.values()]
    last_col = get_column_letter(len(headers))

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = "BÁO CÁO LỖI CHẤT LƯỢNG DỮ LIỆU (DATA QUALITY)"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center")
    meta = [
        f"Phạm vi: {_scope_label(f, user)}",
        f"Bộ lọc: {_describe_filters(f)}",
        f"Người xuất: {user.display_name} ({user.username}) - {datetime.now():%d/%m/%Y %H:%M}",
        f"Tổng số bản ghi: {len(rows):,}" + (f" (đã giới hạn {settings.EXPORT_MAX_ROWS:,} dòng)" if truncated else ""),
    ]
    for i, text in enumerate(meta, start=2):
        ws.merge_cells(f"A{i}:{last_col}{i}")
        ws[f"A{i}"] = text
        ws[f"A{i}"].font = Font(italic=True, color="595959")

    header_row = len(meta) + 3
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=title)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 24

    for idx, row in enumerate(rows, start=1):
        r = header_row + idx
        values = row.model_dump()
        for col, key in enumerate(keys, start=1):
            cell = ws.cell(row=r, column=col, value=idx if key == "STT" else values[key])
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(key == "description"))
            if key in DATE_COLUMNS:
                cell.number_format = "DD/MM/YYYY"
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif key in ("STT", "severity", "branch_code"):
                cell.alignment = Alignment(horizontal="center", vertical="top")
            if key == "severity" and values[key] in SEVERITY_FILL:
                cell.fill = SEVERITY_FILL[values[key]]

    for col, key in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS.get(key, 15)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{last_col}{header_row + max(len(rows), 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"


def _write_summary_sheet(ws2, rows: list[DQErrorOut]) -> None:
    """Sheet tổng hợp: số lỗi theo chi nhánh x mức độ."""
    df = pd.DataFrame([r.model_dump() for r in rows], columns=list(EXPORT_COLUMNS))
    sev_cols = [s.value for s in Severity]
    if df.empty:
        pivot = pd.DataFrame(columns=["branch_code", "branch_name", *sev_cols])
    else:
        pivot = (pd.crosstab([df["branch_code"], df["branch_name"]], df["severity"])
                 .reindex(columns=sev_cols, fill_value=0).reset_index())
    pivot["Tổng"] = pivot[sev_cols].sum(axis=1)
    pivot = pivot.sort_values("Tổng", ascending=False)

    summary_headers = ["Mã chi nhánh", "Tên chi nhánh", *sev_cols, "Tổng"]
    ws2.append(summary_headers)
    for row in pivot.itertuples(index=False):
        ws2.append(list(row))
    total_row = ws2.max_row + 1
    ws2.cell(row=total_row, column=1, value="TỔNG CỘNG")
    for col in range(3, len(summary_headers) + 1):
        letter = get_column_letter(col)
        ws2.cell(row=total_row, column=col, value=f"=SUM({letter}2:{letter}{total_row - 1})")
    for cell in ws2[1]:
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center")
    for row in ws2.iter_rows(min_row=2, max_row=total_row):
        for cell in row:
            cell.border = BORDER
            if cell.row == total_row:
                cell.font = Font(bold=True)
    for letter, width in zip("ABCDEF", (14, 30, 10, 10, 10, 10)):
        ws2.column_dimensions[letter].width = width
    ws2.freeze_panes = "A2"


def _attachment(filename: str) -> dict[str, str]:
    return {"Content-Disposition": f'attachment; filename="{filename}"'}


@router.get("/excel")
def export_excel(user: CurrentUser = Depends(get_current_user), filters: DQFilter = Depends(get_filters),
                 db: Session = Depends(get_db)):
    scoped = user.scope(filters)
    rows, truncated = _load_rows(db, scoped, user, "xlsx")
    buf = build_excel(rows, user, scoped, truncated)
    return StreamingResponse(buf, media_type=XLSX_MIME, headers=_attachment(_filename("xlsx", scoped, user)))


@router.get("/csv")
def export_csv(user: CurrentUser = Depends(get_current_user), filters: DQFilter = Depends(get_filters),
               db: Session = Depends(get_db)):
    scoped = user.scope(filters)
    rows, _ = _load_rows(db, scoped, user, "csv")
    df = pd.DataFrame([r.model_dump() for r in rows], columns=list(EXPORT_COLUMNS))
    for col in DATE_COLUMNS:
        df[col] = pd.to_datetime(df[col])
    df = df.rename(columns=EXPORT_COLUMNS)
    buf = io.BytesIO()
    # utf-8-sig (có BOM) để Excel mở đúng tiếng Việt
    df.to_csv(buf, index=False, encoding="utf-8-sig", date_format="%d/%m/%Y")
    buf.seek(0)
    return StreamingResponse(buf, media_type="text/csv; charset=utf-8",
                             headers=_attachment(_filename("csv", scoped, user)))
