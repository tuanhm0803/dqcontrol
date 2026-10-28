"""File Excel / CSV: yêu cầu loại trừ (file mẫu, đọc file upload, xuất yêu cầu đã duyệt) và danh mục mã quy tắc.

File xuất có các cột cùng tiêu đề với file mẫu, nên có thể sửa rồi upload lại
(cột được nhận theo tiêu đề, cột thừa bị bỏ qua).
"""
import io
import re
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config import settings
from models import EXCLUSION_EXPORT_COLUMNS, CurrentUser, DQRuleOut, ExclusionFilter, ExclusionRequestOut
from routers.export_router import BORDER, HEADER_FILL, HEADER_FONT
from workflow import ExclusionInput, RuleInput

DATA_SHEET = "Yêu cầu loại trừ"
LIST_SHEET = "Danh mục"

# Cột file mẫu: (khóa, tiêu đề, độ rộng). "rule_name" là cột công thức tự hiện tên quy tắc, bị bỏ qua khi đọc file.
TEMPLATE_COLUMNS = [
    ("branch_code", "Mã chi nhánh", 14),
    ("rule_code", "Mã quy tắc DQ *", 18),
    ("rule_name", "Tên quy tắc DQ (tự hiện)", 40),
    ("record_key", "Mã bản ghi *", 22),
    ("reason", "Lý do loại trừ *", 60),
    ("valid_to", "Hiệu lực đến * (dd/mm/yyyy)", 16),
]
TEMPLATE_FORMAT_ROWS = 500  # số dòng được định dạng sẵn / có công thức tên quy tắc
# Tiêu đề cột được chấp nhận khi đọc file (đã chuẩn hóa: chữ thường, bỏ "*" và phần trong ngoặc)
HEADER_ALIASES = {
    "mã chi nhánh": "branch_code", "branch_code": "branch_code",
    "mã quy tắc dq": "rule_code", "mã quy tắc": "rule_code", "mã rule": "rule_code", "rule_code": "rule_code",
    "mã bản ghi": "record_key", "record_key": "record_key",
    "lý do loại trừ": "reason", "lý do": "reason", "reason": "reason",
    "hiệu lực đến": "valid_to", "valid_to": "valid_to",
}
REQUIRED_COLUMNS = {"rule_code": "Mã quy tắc DQ", "record_key": "Mã bản ghi", "reason": "Lý do loại trừ",
                    "valid_to": "Hiệu lực đến"}
EXPORT_WIDTHS = {"branch_code": 13, "rule_code": 16, "rule_name": 40, "record_key": 20, "reason": 50, "valid_to": 13,
                 "id": 11, "requested_by": 14, "requested_at": 17, "manager_by": 14, "manager_at": 17,
                 "business_by": 14, "business_at": 17, "dq_by": 14, "dq_at": 17, "decision_note": 50}
DATE_KEYS = ("valid_to",)
DATETIME_KEYS = ("requested_at", "manager_at", "business_at", "dq_at")

# File danh mục mã quy tắc (admin tải về / upload thay thế toàn bộ)
RULES_SHEET = "Danh mục quy tắc"
RULE_HEADER_ALIASES = {
    "mã quy tắc dq": "rule_code", "mã quy tắc": "rule_code", "mã rule": "rule_code", "rule_code": "rule_code",
    "tên quy tắc dq": "rule_name", "tên quy tắc": "rule_name", "tên rule": "rule_name", "rule_name": "rule_name",
}
RULE_REQUIRED_COLUMNS = {"rule_code": "Mã quy tắc DQ", "rule_name": "Tên quy tắc DQ"}
RULES_UPLOAD_MAX_ROWS = 5000


class UploadError(Exception):
    """File upload không đọc được / sai cấu trúc; message hiển thị cho người dùng."""


# --------------------------------------------------------------------------- #
# File mẫu
# --------------------------------------------------------------------------- #
def _style_header(ws, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def _format_template_rows(ws, letters: dict[str, str]) -> None:
    """Định dạng sẵn các dòng dữ liệu; cột tên quy tắc là công thức tra sheet Danh mục."""
    for key, _, width in TEMPLATE_COLUMNS:
        letter = letters[key]
        ws.column_dimensions[letter].width = width
        # Định dạng Text để Excel không làm mất số 0 đầu / đổi mã dài thành dạng 1E+11
        fmt = "DD/MM/YYYY" if key == "valid_to" else "@"
        for row in range(2, TEMPLATE_FORMAT_ROWS + 1):
            cell = ws[f"{letter}{row}"]
            if key == "rule_name":
                cell.value = f"=IFERROR(VLOOKUP({letters['rule_code']}{row},'{LIST_SHEET}'!$C:$D,2,FALSE),\"\")"
                cell.font = Font(italic=True, color="595959")
            else:
                cell.number_format = fmt


def build_template(branches: list[tuple[str, str]], rules: list[tuple[str, str]], user: CurrentUser) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    ws.append([title for _, title, _ in TEMPLATE_COLUMNS])
    _style_header(ws, len(TEMPLATE_COLUMNS))
    max_row = settings.EXCLUSION_UPLOAD_MAX_ROWS + 1
    letters = {key: get_column_letter(col) for col, (key, _, _) in enumerate(TEMPLATE_COLUMNS, start=1)}
    _format_template_rows(ws, letters)

    lists = wb.create_sheet(LIST_SHEET)
    lists.append(["Mã chi nhánh", "Tên chi nhánh", "Mã quy tắc DQ", "Tên quy tắc DQ"])
    _style_header(lists, 4)
    for i in range(max(len(branches), len(rules))):
        code, name = branches[i] if i < len(branches) else (None, None)
        rule, rule_name = rules[i] if i < len(rules) else (None, None)
        lists.append([code, name, rule, rule_name])
    for letter, width in zip("ABCD", (14, 32, 18, 50)):
        lists.column_dimensions[letter].width = width
    if branches:
        dv = DataValidation(type="list", formula1=f"='{LIST_SHEET}'!$A$2:$A${len(branches) + 1}", allow_blank=True)
        dv.add(f"{letters['branch_code']}2:{letters['branch_code']}{max_row}")
        ws.add_data_validation(dv)
    if rules:
        dv = DataValidation(type="list", formula1=f"='{LIST_SHEET}'!$C$2:$C${len(rules) + 1}", allow_blank=False,
                            showErrorMessage=True, errorTitle="Mã quy tắc", error="Chọn mã quy tắc trong danh mục.")
        dv.add(f"{letters['rule_code']}2:{letters['rule_code']}{max_row}")
        ws.add_data_validation(dv)

    guide = wb.create_sheet("Hướng dẫn")
    branch_note = (f"Bỏ trống = chi nhánh của bạn ({user.branch_code}). " if user.branch_code else "Bắt buộc. ") + \
        "Được khai báo cho chi nhánh khác; chọn trong danh sách (sheet Danh mục)."
    for line in [
        ["HƯỚNG DẪN KHAI BÁO YÊU CẦU LOẠI TRỪ"],
        [],
        ["Mỗi dòng ở sheet '" + DATA_SHEET + "' là một bản ghi cần loại trừ khỏi một quy tắc DQ."],
        ["Không đổi tên dòng tiêu đề. Cột có dấu * là bắt buộc."],
        [],
        ["Cột", "Mô tả"],
        ["Mã chi nhánh", branch_note],
        ["Mã quy tắc DQ *", "Chọn trong danh sách (sheet Danh mục)."],
        ["Tên quy tắc DQ", f"Tự hiện theo mã quy tắc (công thức, {TEMPLATE_FORMAT_ROWS} dòng đầu); không cần nhập."],
        ["Mã bản ghi *", "Số CIF / số tài khoản / số hợp đồng... của bản ghi cần loại trừ (tối đa 100 ký tự)."],
        ["Lý do loại trừ *", f"Tối thiểu {settings.EXCLUSION_REASON_MIN_LENGTH} ký tự."],
        ["Hiệu lực đến *", "Ngày dd/mm/yyyy, từ hôm nay trở đi."],
        [],
        [f"Tối đa {settings.EXCLUSION_UPLOAD_MAX_ROWS:,} dòng, dung lượng {settings.EXCLUSION_UPLOAD_MAX_MB} MB."],
        ["Dòng sai, hoặc trùng yêu cầu đã duyệt còn hiệu lực, sẽ bị bỏ qua; các dòng hợp lệ vẫn được gửi."],
    ]:
        guide.append(line)
    guide["A1"].font = Font(bold=True, size=13, color="1F4E78")
    for cell in guide[6]:
        cell.font = Font(bold=True)
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 95

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# --------------------------------------------------------------------------- #
# Đọc file upload
# --------------------------------------------------------------------------- #
def _norm_header(value) -> str:
    text = re.sub(r"\([^)]*\)", "", str(value or "")).replace("*", "")
    return " ".join(text.split()).lower()


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    return str(value).strip()


def _parse_date(value) -> date | None:
    """Ô ngày: kiểu date của Excel hoặc chuỗi dd/mm/yyyy, yyyy-mm-dd. Raise ValueError nếu sai."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


def _find_header(rows: list[tuple], aliases: dict[str, str], required: dict[str, str],
                 anchor: str) -> tuple[int, dict[str, int]]:
    """Dòng tiêu đề = dòng đầu tiên (trong 10 dòng) có cột `anchor`; trả (chỉ số dòng, khóa -> chỉ số cột)."""
    for idx, row in enumerate(rows[:10]):
        mapping: dict[str, int] = {}
        for col, value in enumerate(row):
            key = aliases.get(_norm_header(value))
            if key and key not in mapping:
                mapping[key] = col
        if anchor in mapping:
            missing = [title for key, title in required.items() if key not in mapping]
            if missing:
                raise UploadError("File thiếu cột: " + ", ".join(missing) + ".")
            return idx, mapping
    raise UploadError(f"Không tìm thấy dòng tiêu đề (cần có cột '{required[anchor]}'). Hãy dùng file mẫu.")


def _read_rows(data: bytes, sheet: str) -> list[tuple]:
    """Các dòng của sheet `sheet` (không có thì sheet đầu tiên) trong file .xlsx."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise UploadError("Không đọc được file. Vui lòng dùng file Excel .xlsx theo mẫu.") from None
    try:
        ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
        return list(ws.iter_rows(values_only=True))
    finally:
        wb.close()


def _data_rows(rows: list[tuple], header_idx: int, cols: dict[str, int], max_rows: int):
    """Duyệt (nhãn 'Dòng n', hàm lấy ô theo khóa) cho các dòng dữ liệu không trống (chỉ xét các cột được đọc,
    bỏ qua cột công thức / cột thừa của file xuất)."""
    count = 0
    for offset, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        def cell(key: str, row=row):
            col = cols.get(key)
            return row[col] if col is not None and col < len(row) else None
        if not any(_text(cell(key)) for key in cols):
            continue
        count += 1
        if count > max_rows:
            raise UploadError(f"File có quá {max_rows:,} dòng dữ liệu. Vui lòng tách thành nhiều file.")
        yield f"Dòng {offset}", cell


def parse_upload(data: bytes) -> tuple[list[ExclusionInput], list[str]]:
    """Đọc file .xlsx -> (các dòng hợp lệ về định dạng, lỗi theo dòng). Raise UploadError nếu cả file hỏng."""
    rows = _read_rows(data, DATA_SHEET)
    header_idx, cols = _find_header(rows, HEADER_ALIASES, REQUIRED_COLUMNS, "record_key")
    items: list[ExclusionInput] = []
    errors: list[str] = []
    for label, cell in _data_rows(rows, header_idx, cols, settings.EXCLUSION_UPLOAD_MAX_ROWS):
        try:
            valid_to = _parse_date(cell("valid_to"))
        except ValueError as exc:
            errors.append(f"{label}: ngày hiệu lực '{exc}' không đúng dạng dd/mm/yyyy")
            continue
        items.append(ExclusionInput(label=label, branch_code=_text(cell("branch_code")) or None,
                                    rule_code=_text(cell("rule_code")), record_key=_text(cell("record_key")),
                                    reason=_text(cell("reason")), valid_to=valid_to))
    if not items and not errors:
        raise UploadError("File không có dòng dữ liệu nào.")
    return items, errors


# --------------------------------------------------------------------------- #
# Danh mục mã quy tắc: tải về (cũng là file mẫu) / đọc file upload
# --------------------------------------------------------------------------- #
def build_rules_excel(rules: list[DQRuleOut]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = RULES_SHEET
    ws.append(["Mã quy tắc DQ *", "Tên quy tắc DQ *", "Trạng thái (bỏ qua khi upload)", "Cập nhật bởi", "Cập nhật lúc"])
    _style_header(ws, 5)
    for r in rules:
        ws.append([r.rule_code, r.rule_name, "Đang dùng" if r.is_active else "Ngừng dùng", r.updated_by, r.updated_at])
        ws.cell(row=ws.max_row, column=5).number_format = "DD/MM/YYYY HH:MM"
    for letter, width in zip("ABCDE", (18, 70, 16, 14, 17)):
        ws.column_dimensions[letter].width = width
    for row in range(2, max(ws.max_row, 200) + 1):
        ws[f"A{row}"].number_format = "@"

    guide = wb.create_sheet("Hướng dẫn")
    for line in [
        ["UPLOAD DANH MỤC MÃ QUY TẮC (THAY THẾ TOÀN BỘ)"],
        [],
        [f"Sheet '{RULES_SHEET}': mỗi dòng một mã quy tắc; cột Mã quy tắc DQ và Tên quy tắc DQ là bắt buộc."],
        ["Mã quy tắc: chữ, số và _ . - (tối đa 50 ký tự, tự chuyển chữ hoa); không được trùng."],
        ["Mọi mã trong file được thêm / cập nhật tên và chuyển sang Đang dùng."],
        ["Mã đang có nhưng KHÔNG có trong file sẽ chuyển sang Ngừng dùng (không khai báo loại trừ mới được, "
         "yêu cầu cũ vẫn hiển thị tên)."],
        ["Chỉ cần một dòng sai thì cả file bị từ chối, danh mục giữ nguyên."],
        ["Các cột Trạng thái / Cập nhật bởi / Cập nhật lúc chỉ để tham khảo, bị bỏ qua khi upload."],
    ]:
        guide.append(line)
    guide["A1"].font = Font(bold=True, size=13, color="1F4E78")
    guide.column_dimensions["A"].width = 120

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def parse_rules_upload(data: bytes) -> list[RuleInput]:
    """Đọc file danh mục -> các dòng (chưa kiểm tra nghiệp vụ; workflow.replace_rules kiểm tra tất cả hoặc không)."""
    rows = _read_rows(data, RULES_SHEET)
    header_idx, cols = _find_header(rows, RULE_HEADER_ALIASES, RULE_REQUIRED_COLUMNS, "rule_code")
    items = [RuleInput(label=label, rule_code=_text(cell("rule_code")), rule_name=_text(cell("rule_name")))
             for label, cell in _data_rows(rows, header_idx, cols, RULES_UPLOAD_MAX_ROWS)]
    if not items:
        raise UploadError("File không có mã quy tắc nào - danh mục giữ nguyên.")
    return items


# --------------------------------------------------------------------------- #
# Xuất yêu cầu loại trừ đã duyệt
# --------------------------------------------------------------------------- #
def describe_filter(f: ExclusionFilter, user: CurrentUser) -> str:
    parts = ["Phạm vi: " + ("Toàn hệ thống" if user.all_branches and not f.branch_code
                            else f"Chi nhánh {f.branch_code or user.branch_code}")]
    if f.rule_code:
        parts.append(f"Mã quy tắc: {f.rule_code}")
    if f.decided_from or f.decided_to:
        start = f.decided_from.strftime("%d/%m/%Y") if f.decided_from else "..."
        end = f.decided_to.strftime("%d/%m/%Y") if f.decided_to else "..."
        parts.append(f"Ngày duyệt: {start} - {end}")
    if f.active_only:
        parts.append("Chỉ yêu cầu còn hiệu lực")
    if f.keyword:
        parts.append(f"Từ khóa: \"{f.keyword}\"")
    return "; ".join(parts)


def build_export_excel(rows: list[ExclusionRequestOut], user: CurrentUser, f: ExclusionFilter,
                       truncated: bool) -> io.BytesIO:
    """Sheet dữ liệu có tiêu đề ở dòng 1 (để hệ thống khác đọc thẳng), sheet 'Thông tin' chứa metadata."""
    keys = list(EXCLUSION_EXPORT_COLUMNS)
    wb = Workbook()
    ws = wb.active
    ws.title = "Loại trừ đã duyệt"
    ws.append(list(EXCLUSION_EXPORT_COLUMNS.values()))
    _style_header(ws, len(keys))
    for r, item in enumerate(rows, start=2):
        values = item.model_dump()
        for c, key in enumerate(keys, start=1):
            cell = ws.cell(row=r, column=c, value=values[key])
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=key in ("reason", "decision_note"))
            if key in DATE_KEYS:
                cell.number_format = "DD/MM/YYYY"
            elif key in DATETIME_KEYS:
                cell.number_format = "DD/MM/YYYY HH:MM"
            elif key == "record_key":
                cell.number_format = "@"
    for c, key in enumerate(keys, start=1):
        ws.column_dimensions[get_column_letter(c)].width = EXPORT_WIDTHS.get(key, 15)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(keys))}{max(len(rows), 1) + 1}"

    info = wb.create_sheet("Thông tin")
    for line in [
        ["DANH SÁCH YÊU CẦU LOẠI TRỪ ĐÃ DUYỆT"],
        [describe_filter(f, user)],
        [f"Người xuất: {user.display_name} ({user.username}) - {datetime.now():%d/%m/%Y %H:%M}"],
        [f"Tổng số bản ghi: {len(rows):,}" + (f" (đã giới hạn {settings.EXPORT_MAX_ROWS:,} dòng)" if truncated else "")],
    ]:
        info.append(line)
    info["A1"].font = Font(bold=True, size=13, color="1F4E78")
    info.column_dimensions["A"].width = 100

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_export_csv(rows: list[ExclusionRequestOut]) -> io.BytesIO:
    df = pd.DataFrame([r.model_dump() for r in rows], columns=list(EXCLUSION_EXPORT_COLUMNS))
    df["valid_to"] = pd.to_datetime(df["valid_to"]).dt.strftime("%d/%m/%Y")
    for key in DATETIME_KEYS:
        df[key] = pd.to_datetime(df[key]).dt.strftime("%d/%m/%Y %H:%M")
    # Mỗi yêu cầu đúng 1 dòng CSV (nhật ký ghi chú nhiều dòng -> nối bằng " | ") cho hệ thống đọc theo dòng
    for key in ("reason", "decision_note"):
        df[key] = df[key].fillna("").str.replace(r"\s*\n\s*", " | ", regex=True)
    buf = io.BytesIO()
    # utf-8-sig (có BOM) để Excel mở đúng tiếng Việt
    df.rename(columns=EXCLUSION_EXPORT_COLUMNS).to_csv(buf, index=False, encoding="utf-8-sig")
    buf.seek(0)
    return buf
